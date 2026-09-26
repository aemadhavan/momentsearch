"""Verification test suite for US-03: Slide Deck Ingestion Pipeline."""
from __future__ import annotations

import os
import sys
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches

from src import config, db
from src.ingest.deck import ingest_deck, t_parse_chunk_deck
from src.rag import vector_store
from src.rag.embeddings import embed_query


def create_sample_pdf_deck(file_path: Path) -> None:
    """Create a 3-slide PDF deck."""
    doc = fitz.open()

    # Slide 1
    s1 = doc.new_page(width=720, height=405)  # 16:9
    s1.insert_text(
        (50, 80),
        "ARGUS: Moment Search at Scale\n\n"
        "Keynote Presentation — Scalable Multimodal Knowledge Retrieval\n"
        "Presented by the Forward Deployed Engineering Team",
        fontsize=16,
    )

    # Slide 2: The exact slide checked in eval.py ("the slide about one index for every source")
    s2 = doc.new_page(width=720, height=405)
    s2.insert_text(
        (50, 80),
        "One Index for Every Source\n\n"
        "Unified hybrid retrieval across video moments, academic papers, and conference decks.\n"
        "One index for every source: a single shared Qdrant collection stores embeddings for all media types with "
        "modality-specific locators (start_ms, page, slide).",
        fontsize=16,
    )

    # Slide 3: Diagrammatic slide
    s3 = doc.new_page(width=720, height=405)
    s3.insert_text(
        (50, 80),
        "System Architecture & Prefect Queue Decoupling\n\n"
        "Fast-path Admin API (HTTP 202) -> Prefect Cloud Queue -> Async Workers -> Qdrant Cloud.",
        fontsize=16,
    )

    doc.save(str(file_path))
    doc.close()


def create_sample_pptx_deck(file_path: Path) -> None:
    """Create a 2-slide PPTX deck to verify python-pptx extraction."""
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]  # Title & content
    
    # Slide 1
    slide1 = prs.slides.add_slide(slide1_layout := prs.slide_layouts[0])
    slide1.shapes.title.text = "PPTX Deck Ingestion Test"
    slide1.placeholders[1].text = "Testing native PPTX parsing in ARGUS."

    # Slide 2
    slide2 = prs.slides.add_slide(slide_layout)
    slide2.shapes.title.text = "Decoupled Queue Scaling"
    slide2.placeholders[1].text = "Workers consume ingestion jobs without impacting concurrent search latencies."

    prs.save(str(file_path))


def test_deck_ingestion():
    test_dir = ROOT / "data" / "test_media"
    test_dir.mkdir(parents=True, exist_ok=True)
    
    # 1. Test PDF slide deck
    pdf_deck_path = test_dir / "sample_keynote.pdf"
    create_sample_pdf_deck(pdf_deck_path)
    print(f"[1] Created test PDF deck at: {pdf_deck_path}")

    doc_id = f"doc_{uuid.uuid4().hex[:8]}"
    user_id = config.SINGLE_USER_ID

    # Parse chunks directly
    chunks = t_parse_chunk_deck.fn(doc_id, user_id, str(pdf_deck_path))
    assert len(chunks) == 3, f"Expected 3 slide chunks, got {len(chunks)}"
    slides_found = [c.slide for c in chunks]
    assert slides_found == [1, 2, 3], f"Expected slides [1, 2, 3], got {slides_found}"
    print(f"[2] PDF slide extraction verified: {len(chunks)} slides {slides_found}")

    # Register in DB
    db.create_document(
        doc_id=doc_id,
        user_id=user_id,
        kind="deck",
        uri=str(pdf_deck_path),
        title="KDD Keynote Deck",
    )
    row = db.get_document(doc_id)
    assert row is not None and row["status"] == "pending"
    print(f"[3] Inserted pending deck {doc_id} in DB.")

    # Run ingestion flow
    result = ingest_deck.fn(doc_id, user_id)
    print(f"[4] Ingest deck flow result: {result}")
    assert result["kind"] == "deck"
    assert result["chunks"] == 3

    # Check DB status
    updated_row = db.get_document(doc_id)
    assert updated_row["status"] == "indexed"
    assert updated_row["chunk_count"] == 3
    assert updated_row["progress"] == 1.0
    print(f"[5] DB status verified: {updated_row['status']} ({updated_row['chunk_count']} slides)")

    # Search for "the slide about one index for every source" (exact eval.py query)
    query_vec = embed_query("the slide about one index for every source")
    hits = vector_store.search_text(query_vec, user_id, top_k=5)
    assert len(hits) > 0, "No hits returned for slide query"
    
    deck_hits = [h for h in hits if h.get("source_id") == doc_id or h.get("doc_id") == doc_id]
    assert len(deck_hits) > 0, "Ingested deck was not retrieved in top search hits"
    
    for i, h in enumerate(deck_hits):
        print(f"[hit {i}] score={h.get('score'):.4f}, locator={h.get('locator')}, text={h.get('text', '')[:60]}...")
    top_hit = deck_hits[0]
    print(f"[6] Top retrieved deck hit: score={top_hit.get('score'):.4f}, locator={top_hit.get('locator')}")
    assert top_hit.get("kind") == "deck", f"Expected kind 'deck', got {top_hit.get('kind')}"
    assert "locator" in top_hit, "locator missing from payload"
    assert top_hit["locator"].get("slide") == 2, f"Expected slide 2 for 'one index for every source', got {top_hit['locator']}"
    print(f"[7] Exact slide locator verified: Slide {top_hit['locator']['slide']} matched perfectly!")

    # 2. Test PPTX format parsing
    pptx_path = test_dir / "sample_deck.pptx"
    create_sample_pptx_deck(pptx_path)
    pptx_doc_id = f"doc_{uuid.uuid4().hex[:8]}"
    pptx_chunks = t_parse_chunk_deck.fn(pptx_doc_id, user_id, str(pptx_path))
    assert len(pptx_chunks) == 2, f"Expected 2 slides from PPTX, got {len(pptx_chunks)}"
    assert [c.slide for c in pptx_chunks] == [1, 2]
    print(f"[8] PPTX slide extraction verified: {len(pptx_chunks)} slides parsed successfully!")


if __name__ == "__main__":
    test_deck_ingestion()
    print("\nALL US-03 ACCEPTANCE CRITERIA MET!")
